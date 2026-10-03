from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r'G:\project\TapTapGJ21\涌现_团队执行版.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Microsoft YaHei'
    doc.styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
doc.styles['Normal'].font.size = Pt(10.5)

def shade(cell, fill='4472C4'):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def make_table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        shade(c)
        for run in c.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    doc.add_paragraph()

def bullet(text):
    doc.add_paragraph(text, style='List Bullet')

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('涌现 团队执行版策划案')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('目标：用最少规则验证“拆墙 → 创造 → 记忆显现 → 新问题”的核心体验').italic = True

doc.add_heading('一 项目一句话', 1)
doc.add_paragraph('玩家拆分、移动、旋转和融合墙体，创造出满足某种行为意义的结构；系统通过环境反馈和记忆片段告诉玩家这里曾经发生过什么，并提出下一个问题。')

doc.add_heading('二 核心循环', 1)
doc.add_paragraph('玩家循环：分裂 → 重组 → 膨胀 → 再分裂。叙事循环：空白 → 记忆显现 → 意义理解 → 新问题 → 新空白。两者必须在同一段可玩流程中发生。')
make_table(['步骤', '玩家做什么', '系统判断', '玩家看到什么'], [
    ('1 发现', '尝试拆下墙体', '墙是否可交互', '墙体松动、裂纹、声音反馈'),
    ('2 分裂', '把墙拆成多个物块', '记录物块 ID、位置、状态和完整度', '分裂动画与可操作物块'),
    ('3 重组', '移动、旋转、颠倒、融合', '计算结构关系和空间条件', '平台、开口、围合等结构结果'),
    ('4 膨胀', '加入形态、颜色、重量或反射', '检查环境互动是否产生新现象', '光路、倾斜、通路、承托等结果'),
    ('5 意图确认', '坐下、躺下、停留、穿过', '结构约束 + 语义倾向 + 行为确认 + 上下文', '记忆残影、声音或墙面变化'),
    ('6 新问题', '继续探索下一处空白', '更新空间、章节状态和记忆', '新的线索和未解释的问题'),
])

doc.add_heading('三 机制范围', 1)
make_table(['层级', '首版要做', '暂不做'], [
    ('结构', '分裂、移动、旋转/颠倒、融合', '复杂破坏、程序生成建筑'),
    ('形态', '3–5 种基础几何变体', '大量家具图鉴'),
    ('状态', '完整度、颜色', '复杂材料库'),
    ('属性', '重量、反射', '穿透、粘附等后续扩展'),
    ('语义', '休息、通行、等待，首版至少做 2 个', '所有物品自动识别'),
    ('反馈', '裂纹、残影、声音、墙面记忆', '长篇文字解释'),
])

doc.add_heading('四 语义判定规则', 1)
doc.add_paragraph('系统不判断“玩家是否造出了标准床”，而判断“当前结构和行为是否足以唤醒休息记忆”。')
make_table(['字段', '执行定义'], [
    ('必要条件', '缺少就不能成立，例如存在承托面、角色可到达'),
    ('支持条件', '提高成立倾向，例如靠墙、半封闭、柔软'),
    ('确认行为', '玩家主动表达意图，例如坐下、躺下、停留'),
    ('叙事上下文', '当前章节和空间是否允许该记忆出现'),
    ('反馈', '视觉、声音、镜头、墙面状态变化'),
    ('下一问题', '本次记忆显现后玩家要继续追问什么'),
])
doc.add_paragraph('判定顺序：先过滤不可能结构，再计算语义倾向，最后用玩家行为确认。禁止仅凭外形相似度直接触发剧情。')

doc.add_heading('五 首个垂直切片', 1)
make_table(['阶段', '内容', '完成标准'], [
    ('P0 空白房间', '1 个空白房间、4 面墙、基础输入', '玩家能发现至少 1 面墙可拆'),
    ('P1 结构变换', '分裂、移动、旋转、融合、重置', '玩家能独立完成一次重组'),
    ('P2 第一次涌现', '承托结构 + 休息节点', '玩家用非配方方式触发记忆'),
    ('P3 形态状态', '形态装置、颜色/完整度反馈', '同一批物块能产生至少 2 种不同结果'),
    ('P4 通行与收束', '通行节点、相邻空间、新问题', '完成“发现规则→创造→理解记忆→产生新问题”闭环'),
])

doc.add_heading('六 三职能工作清单', 1)
make_table(['职能', '本阶段必须交付', '验收方式'], [
    ('策划', '机制规则、语义节点表、测试案例、关卡灰盒、记忆事件表', '能写清成立/不成立条件，并列出预期意外结果'),
    ('程序', 'Wall/Block、Transform、Shape/State/Property、Semantic Evaluator、Memory、Debug 面板', '系统可配置、可重置、可解释“为什么触发”'),
    ('美术', '墙体状态、分裂动画、物块材质、装置表现、记忆 VFX/声音规范', '玩家无需长文字即可理解可操作状态和因果方向'),
])

doc.add_heading('七 每轮迭代流程', 1)
for text in [
    '定义：策划写清玩家意图、规则、边界和预期意外结果。',
    '灰盒：程序用占位物实现，策划提供最小测试空间。',
    '联调：三方试玩，调整阈值、手感、反馈和语义判定。',
    '验证：确认结果来自规则组合，而不是单个脚本。',
    '固化：把验证成功的规则抽象成数据配置和可复用资产。',
    '扩展：只增加能与至少两个旧机制组合的新内容。',
]:
    bullet(text)

doc.add_heading('八 验收标准', 1)
for text in [
    '玩家不依赖长篇文字即可发现墙可以被拆分。',
    '分裂、移动、旋转、融合足以产生“我在创造”的感觉。',
    '至少出现一种策划没有逐项写死、但可以重复发生的组合结果。',
    '同一语义节点允许多种表达，但不会因“看起来像”而随机触发。',
    '玩家偏离主线时，世界通过空间线索、残影或新问题吸引其回归。',
    '完整度损耗会影响记忆反馈，但不会造成关键流程软锁。',
    '策划可以通过配置修改阈值、权重、反馈和下一问题，无需为单个关卡硬编码。',
]:
    bullet(text)

doc.add_heading('九 当前版本的取舍', 1)
doc.add_paragraph('首版不做传统 Crafting 配方、大量材料背包、技能树等级、大量独立机关、固定家具收集和完全开放的语义识别。所有新增需求先回答三个问题：是否增加组合可能？是否接入现有规则链？是否帮助玩家感受到涌现或推进主线？若都不能肯定，暂缓。')

doc.add_heading('十 首周建议任务', 1)
make_table(['负责人', '任务', '产出'], [
    ('策划', '完成分裂/移动/旋转/融合规则卡；定义“休息”和“通行”节点', '机制卡、语义节点表、10 个测试案例'),
    ('程序', '完成墙体拆分、物块操作、状态继承、重置和 Debug 日志', '可玩灰盒版本'),
    ('美术', '确定白色空房间、墙体裂纹、分裂动画和第一段记忆反馈方向', '占位视觉规范 + 反馈样例'),
    ('共同', '每天用同一测试房试玩并记录意外结果', '问题清单、参数变更记录、次日验证目标'),
])
doc.add_paragraph('版本出口：团队可以实际玩到“拆墙 → 组合 → 产生非预设结果 → 触发记忆 → 留下新问题”，并能解释每一步为什么发生。')
doc.save(OUT)
print(OUT)
