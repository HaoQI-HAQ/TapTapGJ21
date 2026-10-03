from docx import Document

p = r'D:\桌面\.._游戏策划案_核心循环与涌现叙事_0.3.docx'
d = Document(p)
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if t:
        print(f'[P{i}] {t}')
for ti, table in enumerate(d.tables):
    print(f'\n[TABLE {ti}]')
    for row in table.rows:
        print(' | '.join(c.text.replace('\n', ' / ') for c in row.cells))
